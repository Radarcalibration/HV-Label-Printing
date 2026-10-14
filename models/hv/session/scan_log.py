from __future__ import annotations

import os
import sqlite3
from contextlib import contextmanager
from dataclasses import dataclass
from typing import Iterator

import openpyxl

from models.hv.session.paths import get_base_dir

LOG_FILE = get_base_dir() / "data" / "scan_log.xlsx"
DB_FILE = get_base_dir() / "data" / "scan_log.db"

SHEET_NAME = "HV Dummy"

COLUMNS = [
    "Sl.No",
    "Date",
    "Time",
    "Serial Count",
    "Pack ID (Serial Number)",
    "Ref No",
    "Variant",
    "M1",
    "M2",
    "M3",
    "M4",
    "Dummy Pack ID Status",
    "Dummy Pack QR Data",
    "Issue Description 1",
    "Issue Description 2",
    "Action Plan",
    "Remark",
    "Admin Acknowledgement",
]

_MIGRATION_NAME = "scan_log_xlsx_v1"
_RECORD_FIELDS = (
    "date",
    "time",
    "serial_count",
    "ref_no",
    "variant",
    "m1",
    "m2",
    "m3",
    "m4",
    "dummy_pack_status",
    "dummy_pack_qr",
    "serial_number",
    "issue_desc_1",
    "issue_desc_2",
    "action_plan",
    "remark",
    "admin_ack",
)
_COLUMN_FIELDS = {
    "Date": "date",
    "Time": "time",
    "Serial Count": "serial_count",
    "Ref No": "ref_no",
    "Variant": "variant",
    "M1": "m1",
    "M2": "m2",
    "M3": "m3",
    "M4": "m4",
    "Dummy Pack ID Status": "dummy_pack_status",
    "Dummy Pack QR Data": "dummy_pack_qr",
    "Pack ID (Serial Number)": "serial_number",
    "Issue Description 1": "issue_desc_1",
    "Issue Description 2": "issue_desc_2",
    "Action Plan": "action_plan",
    "Remark": "remark",
    "Admin Acknowledgement": "admin_ack",
}


@dataclass
class ScanRecord:
    date: str = ""
    time: str = ""
    serial_count: str = ""
    ref_no: str = ""
    variant: str = ""
    m1: str = ""
    m2: str = ""
    m3: str = ""
    m4: str = ""
    dummy_pack_status: str = ""
    dummy_pack_qr: str = ""
    serial_number: str = ""
    issue_desc_1: str = ""
    issue_desc_2: str = ""
    action_plan: str = ""
    remark: str = ""
    admin_ack: str = ""

    @property
    def has_issue(self) -> bool:
        return bool(self.issue_desc_1.strip() or self.issue_desc_2.strip())

    @property
    def is_unacknowledged(self) -> bool:
        return self.has_issue and not self.admin_ack.strip()

    def as_row(self, sl_no: int) -> list[str | int]:
        return [
            sl_no,
            self.date,
            self.time,
            self.serial_count,
            self.serial_number,
            self.ref_no,
            self.variant,
            self.m1,
            self.m2,
            self.m3,
            self.m4,
            self.dummy_pack_status,
            self.dummy_pack_qr,
            self.issue_desc_1,
            self.issue_desc_2,
            self.action_plan,
            self.remark,
            self.admin_ack,
        ]


@contextmanager
def _connect_db() -> Iterator[sqlite3.Connection]:
    DB_FILE.parent.mkdir(parents=True, exist_ok=True)
    connection = sqlite3.connect(DB_FILE, timeout=10)
    connection.row_factory = sqlite3.Row
    try:
        connection.execute("PRAGMA busy_timeout = 10000")
        connection.execute(
            """
            CREATE TABLE IF NOT EXISTS scan_records (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                date TEXT NOT NULL DEFAULT '',
                time TEXT NOT NULL DEFAULT '',
                serial_count TEXT NOT NULL DEFAULT '',
                ref_no TEXT NOT NULL DEFAULT '',
                variant TEXT NOT NULL DEFAULT '',
                m1 TEXT NOT NULL DEFAULT '',
                m2 TEXT NOT NULL DEFAULT '',
                m3 TEXT NOT NULL DEFAULT '',
                m4 TEXT NOT NULL DEFAULT '',
                dummy_pack_status TEXT NOT NULL DEFAULT '',
                dummy_pack_qr TEXT NOT NULL DEFAULT '',
                serial_number TEXT NOT NULL DEFAULT '',
                issue_desc_1 TEXT NOT NULL DEFAULT '',
                issue_desc_2 TEXT NOT NULL DEFAULT '',
                action_plan TEXT NOT NULL DEFAULT '',
                remark TEXT NOT NULL DEFAULT '',
                admin_ack TEXT NOT NULL DEFAULT ''
            )
            """
        )
        connection.execute(
            "CREATE TABLE IF NOT EXISTS scan_log_migrations "
            "(name TEXT PRIMARY KEY)"
        )
        connection.execute("BEGIN IMMEDIATE")
        _migrate_xlsx(connection)
        yield connection
        connection.commit()
    except BaseException:
        connection.rollback()
        raise
    finally:
        connection.close()


def _migrate_xlsx(connection: sqlite3.Connection) -> None:
    if connection.execute(
        "SELECT 1 FROM scan_log_migrations WHERE name = ?", (_MIGRATION_NAME,)
    ).fetchone():
        return

    if not LOG_FILE.exists():
        connection.execute(
            "INSERT INTO scan_log_migrations(name) VALUES (?)", (_MIGRATION_NAME,)
        )
        return

    try:
        workbook = openpyxl.load_workbook(LOG_FILE, data_only=True, read_only=True)
    except (OSError, ValueError, openpyxl.utils.exceptions.InvalidFileException) as exc:
        raise RuntimeError(f"Could not import scan history from {LOG_FILE}: {exc}") from exc

    try:
        if SHEET_NAME in workbook.sheetnames:
            worksheet = workbook[SHEET_NAME]
            header = next(
                worksheet.iter_rows(min_row=1, max_row=1, values_only=True), ()
            )
            columns = {
                str(value).strip(): index
                for index, value in enumerate(header)
                if value is not None
            }
            insert_sql = (
                "INSERT INTO scan_records ("
                + ", ".join(_RECORD_FIELDS)
                + ") VALUES ("
                + ", ".join("?" for _ in _RECORD_FIELDS)
                + ")"
            )
            rows = []
            for row in worksheet.iter_rows(min_row=2, values_only=True):
                if not any(value is not None and str(value).strip() for value in row):
                    continue
                record_values = []
                for column_name, field_name in _COLUMN_FIELDS.items():
                    column_index = columns.get(column_name)
                    value = (
                        row[column_index]
                        if column_index is not None and column_index < len(row)
                        else ""
                    )
                    record_values.append("" if value is None else str(value))
                rows.append(record_values)
            if rows:
                connection.executemany(insert_sql, rows)
    finally:
        workbook.close()

    connection.execute(
        "INSERT INTO scan_log_migrations(name) VALUES (?)", (_MIGRATION_NAME,)
    )


def append_scan_record(record: ScanRecord, *, excel_password: str = "06082003") -> None:
    values = [getattr(record, field_name) for field_name in _RECORD_FIELDS]
    placeholders = ", ".join("?" for _ in _RECORD_FIELDS)
    with _connect_db() as connection:
        connection.execute(
            f"INSERT INTO scan_records ({', '.join(_RECORD_FIELDS)}) "
            f"VALUES ({placeholders})",
            values,
        )


def load_all_records_with_rows() -> list[tuple[int, ScanRecord]]:
    with _connect_db() as connection:
        rows = connection.execute(
            f"SELECT id, {', '.join(_RECORD_FIELDS)} "
            "FROM scan_records ORDER BY id"
        ).fetchall()
    return [
        (
            row["id"],
            ScanRecord(**{field_name: row[field_name] for field_name in _RECORD_FIELDS}),
        )
        for row in rows
    ]


def load_all_records() -> list[ScanRecord]:
    return [record for _record_id, record in load_all_records_with_rows()]


def load_unacknowledged_issues() -> list[tuple[int, ScanRecord]]:
    return [
        (record_id, record)
        for record_id, record in load_all_records_with_rows()
        if record.is_unacknowledged
    ]


def patch_admin_ack(
    record_id: int, ack_text: str, *, excel_password: str = "06082003"
) -> bool:
    with _connect_db() as connection:
        cursor = connection.execute(
            "UPDATE scan_records SET admin_ack = ? WHERE id = ?",
            (ack_text, record_id),
        )
        return cursor.rowcount > 0


def open_log_file() -> None:
    if LOG_FILE.exists():
        os.startfile(str(LOG_FILE))
