from __future__ import annotations

from pathlib import Path
from typing import Any

import openpyxl

from models.hv.session.bms_scan_log import LOG_FILE as BMS_LOG_FILE
from models.hv.session.bms_scan_log import SHEET_NAME as BMS_SHEET
from models.hv.session.mongo_store import insert_legacy_record
from models.hv.session.pack_id_scan_log import LOG_FILE as PACK_ID_LOG_FILE
from models.hv.session.pack_id_scan_log import SHEET_NAME as PACK_ID_SHEET
from models.hv.session.scan_log import LOG_FILE as DUMMY_LOG_FILE
from models.hv.session.scan_log import SHEET_NAME as DUMMY_SHEET

_MIGRATIONS = (
    (
        "dummy_scan_logs",
        DUMMY_LOG_FILE,
        DUMMY_SHEET,
        {
            "Date": "date",
            "Time": "time",
            "Serial Count": "serial_count",
            "Pack ID (Serial Number)": "serial_number",
            "Ref No": "ref_no",
            "Variant": "variant",
            "M1": "m1",
            "M2": "m2",
            "M3": "m3",
            "M4": "m4",
            "Dummy Pack ID Status": "dummy_pack_status",
            "Dummy Pack QR Data": "dummy_pack_qr",
            "Issue Description 1": "issue_desc_1",
            "Issue Description 2": "issue_desc_2",
            "Action Plan": "action_plan",
            "Remark": "remark",
            "Admin Acknowledgement": "admin_ack",
        },
    ),
    (
        "pack_id_scan_logs",
        PACK_ID_LOG_FILE,
        PACK_ID_SHEET,
        {
            "Date": "date",
            "Time": "time",
            "Serial Number": "serial_number",
            "Variant": "variant",
            "Pack QR Data": "pack_qr_data",
        },
    ),
    (
        "bms_scan_logs",
        BMS_LOG_FILE,
        BMS_SHEET,
        {
            "Date": "date",
            "Time": "time",
            "Pack QR Data": "pack_qr_data",
            "BMB ID": "bmb_id",
            "CMB ID": "cmb_id",
            "Rework BMB ID": "rework_bmb_id",
            "Rework CMB ID": "rework_cmb_id",
        },
    ),
)


def _migrate_workbook(
    collection_name: str,
    path: Path,
    sheet_name: str,
    field_map: dict[str, str],
) -> tuple[int, int]:
    if not path.exists():
        return 0, 0

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            return 0, 0
        worksheet = workbook[sheet_name]
        headers = [
            str(value).strip() if value is not None else ""
            for value in next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True))
        ]
        imported = skipped = 0
        for row_number, values in enumerate(
            worksheet.iter_rows(min_row=2, values_only=True), start=2
        ):
            if not any(value is not None and str(value).strip() for value in values):
                continue
            row: dict[str, Any] = dict(zip(headers, values))
            record = {
                field: str(row.get(column) or "")
                for column, field in field_map.items()
            }
            if insert_legacy_record(
                collection_name,
                record,
                source_file=str(path.resolve()),
                row_number=row_number,
            ):
                imported += 1
            else:
                skipped += 1
        return imported, skipped
    finally:
        workbook.close()


def main() -> None:
    total_imported = total_skipped = 0
    for collection_name, path, sheet_name, field_map in _MIGRATIONS:
        imported, skipped = _migrate_workbook(
            collection_name, path, sheet_name, field_map
        )
        print(f"{path.name}: imported {imported}; already migrated {skipped}")
        total_imported += imported
        total_skipped += skipped
    print(f"Migration complete: {total_imported} imported; {total_skipped} skipped.")


if __name__ == "__main__":
    main()
