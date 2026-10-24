 

import os
import stat
from dataclasses import dataclass
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.protection import SheetProtection

from models.hv.session.paths import get_base_dir

LOG_FILE = get_base_dir() / "data" / "bms_scan_log.xlsx"
SHEET_NAME = "bmb_cmb"

PROTECT_SHEET = False

COLUMNS = [
    "Sl.No", "Date", "Time", "Pack QR Data",
    "BMB ID", "CMB ID", "Status", "Machine Name", "Machine ID",
]

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(name="Segoe UI", bold=True, color="FFFFFF", size=10)
CELL_FONT   = Font(name="Segoe UI", size=10)
ALT_FILL    = PatternFill("solid", fgColor="EEF2F7")

COL_WIDTHS = {
    "Sl.No":           7,
    "Date":           13,
    "Time":           12,
    "Pack QR Data":   40,
    "BMB ID":         26,
    "CMB ID":         26,
    "Status":         14,
    "Machine Name":   22,
    "Machine ID":     18,
}

LEFT_ALIGN_COLS = {"Pack QR Data", "BMB ID", "CMB ID"}
REWORK_FONT = Font(name="Segoe UI", size=10, color="FF0000FF")


@dataclass
class BmsScanRecord:
    date: str          = ""
    time: str          = ""
    pack_qr_data: str  = ""
    bmb_id: str        = ""
    cmb_id: str        = ""
    status: str        = "Original"
    machine_name: str  = ""
    machine_id: str    = ""
    reworked_kind: str = ""

    def as_row(self, sl_no: int) -> list:
        return [
            sl_no, self.date, self.time, self.pack_qr_data,
            self.bmb_id, self.cmb_id, self.status, self.machine_name, self.machine_id,
        ]


def _ensure_writable(path: Path) -> None:
    try:
        path.parent.mkdir(parents=True, exist_ok=True)
        path.parent.chmod(stat.S_IRWXU | stat.S_IRWXG | stat.S_IRWXO)
        if path.exists():
            path.chmod(
                stat.S_IRUSR | stat.S_IWUSR |
                stat.S_IRGRP | stat.S_IWGRP |
                stat.S_IROTH | stat.S_IWOTH
            )
    except OSError:
        pass


def _protection(password: str, enabled: bool = True) -> SheetProtection:
    if not enabled:
        return SheetProtection(sheet=False)
    return SheetProtection(
        sheet=True, password=password,
        selectLockedCells=True, selectUnlockedCells=True,
        insertRows=False, insertColumns=False, deleteRows=False,
        deleteColumns=False, sort=False, autoFilter=False,
    ) #tambi


def _build_header(ws, excel_password: str) -> None:
    for col_idx, col_name in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col_idx)].width = COL_WIDTHS.get(col_name, 18)
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 28
    ws.protection = _protection(excel_password, enabled=PROTECT_SHEET)


def _get_sheet(wb: openpyxl.Workbook, excel_password: str):
    if SHEET_NAME in wb.sheetnames:
        return wb[SHEET_NAME]
    ws = wb.create_sheet(title=SHEET_NAME, index=0)
    _build_header(ws, excel_password)
    return ws


def _record_from_values(values: dict[str, str]) -> tuple[BmsScanRecord, str]:
    legacy_bmb_rework = values.get("Rework BMB ID", "").strip()
    legacy_cmb_rework = values.get("Rework CMB ID", "").strip()
    reworked_kind = ""
    if legacy_bmb_rework:
        reworked_kind = "BMB"
    if legacy_cmb_rework:
        reworked_kind = "CMB" if not reworked_kind else "BOTH"

    status = values.get("Status", "").strip()
    if not status:
        status = "Rework 1" if reworked_kind else "Original"

    record = BmsScanRecord(
        date=values.get("Date", ""),
        time=values.get("Time", ""),
        pack_qr_data=values.get("Pack QR Data", ""),
        bmb_id=legacy_bmb_rework or values.get("BMB ID", ""),
        cmb_id=legacy_cmb_rework or values.get("CMB ID", ""),
        status=status,
        machine_name=values.get("Machine Name", ""),
        machine_id=values.get("Machine ID", ""),
        reworked_kind=reworked_kind,
    )
    return record, reworked_kind


def _write_data_row(ws, row_idx: int, record: BmsScanRecord, sl_no: int) -> None:
    for col_idx, (col_name, value) in enumerate(zip(COLUMNS, record.as_row(sl_no)), start=1):
        cell = ws.cell(row=row_idx, column=col_idx, value=value)
        is_reworked_id = col_name in {"BMB ID", "CMB ID"} and (
            record.reworked_kind == col_name.removesuffix(" ID")
            or record.reworked_kind == "BOTH"
        )
        cell.font = REWORK_FONT if is_reworked_id else CELL_FONT
        h_align = "left" if col_name in LEFT_ALIGN_COLS else "center"
        cell.alignment = Alignment(horizontal=h_align, vertical="center", wrap_text=True)
        if row_idx % 2 == 0:
            cell.fill = ALT_FILL
    ws.row_dimensions[row_idx].height = 18


def _ensure_workbook(excel_password: str) -> openpyxl.Workbook:
    if LOG_FILE.exists():
        wb = openpyxl.load_workbook(LOG_FILE)
        ws = _get_sheet(wb, excel_password)
        existing_headers = [
            str(ws.cell(row=1, column=i).value or "").strip()
            for i in range(1, ws.max_column + 1)
        ]
        if existing_headers != COLUMNS:
            old_headers = existing_headers
            old_max_column = ws.max_column
            migrated = []
            for row in ws.iter_rows(min_row=2, values_only=True):
                values = {
                    header: str(row[index] or "")
                    for index, header in enumerate(old_headers)
                    if index < len(row) and header
                }
                if not any(values.values()):
                    continue
                record, reworked_kind = _record_from_values(values)
                if reworked_kind:
                    migrated.append(BmsScanRecord(
                        date=record.date,
                        time=record.time,
                        pack_qr_data=record.pack_qr_data,
                        bmb_id=values.get("BMB ID", ""),
                        cmb_id=values.get("CMB ID", ""),
                        status="Original",
                        machine_name=record.machine_name,
                        machine_id=record.machine_id,
                    ))
                    record.status = values.get("Status", "").strip() or "Rework 1"
                    migrated.append(record)
                else:
                    record.status = record.status or "Original"
                    migrated.append(record)
            ws.delete_rows(1, ws.max_row)
            if old_max_column:
                ws.delete_cols(1, old_max_column)
            _build_header(ws, excel_password)
            for index, record in enumerate(migrated, start=2):
                _write_data_row(ws, index, record, index - 1)
        _ensure_writable(LOG_FILE)
        wb.save(LOG_FILE)
        return wb
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = SHEET_NAME
    _build_header(ws, excel_password)
    _ensure_writable(LOG_FILE)
    wb.save(LOG_FILE)
    return wb


def append_scan_record(record: BmsScanRecord, *, excel_password: str = "06082003") -> None:
    wb = _ensure_workbook(excel_password)
    ws = _get_sheet(wb, excel_password)
    ws.protection = _protection(excel_password, enabled=False)

    row_idx = ws.max_row + 1
    _write_data_row(ws, row_idx, record, row_idx - 1)
    ws.protection = _protection(excel_password, enabled=PROTECT_SHEET)
    _ensure_writable(LOG_FILE)
    wb.save(LOG_FILE)

def load_all_records() -> list[BmsScanRecord]:
    return [rec for _row, rec in load_all_records_with_rows()]


def load_all_records_with_rows() -> list[tuple[int, BmsScanRecord]]:
    if not LOG_FILE.exists():
        return []
    wb = openpyxl.load_workbook(LOG_FILE, data_only=True)
    try:
        if SHEET_NAME not in wb.sheetnames:
            return []
        ws = wb[SHEET_NAME]
        headers = [
            str(ws.cell(row=1, column=i).value or "").strip()
            for i in range(1, ws.max_column + 1)
        ]
        records: list[tuple[int, BmsScanRecord]] = []
        for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            if not any(row):
                continue
            values = {
                header: str(row[index] or "")
                for index, header in enumerate(headers)
                if index < len(row) and header
            }
            records.append((row_idx, _record_from_values(values)[0]))
        return records
    finally:
        wb.close()


def find_record_by_pack_qr(pack_qr_data: str) -> BmsScanRecord | None:
   
    if not pack_qr_data:
        return None
    target = pack_qr_data.strip()
    for _row, rec in reversed(load_all_records_with_rows()):
        if rec.pack_qr_data.strip() == target:
            return rec
    return None


def open_log_file() -> None:
    if LOG_FILE.exists():
        wb = _ensure_workbook("06082003")
        wb.close()
        os.startfile(str(LOG_FILE))